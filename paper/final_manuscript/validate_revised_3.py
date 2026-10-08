"""Deterministic QA for the revised-3 article and targeted Gate8 package."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
MD = HERE / "SVM_FINAL_PAPER_REVISED_3.md"
DOCX = HERE / "SVM_FINAL_PAPER_REVISED_3.docx"
PDF = HERE / "SVM_FINAL_PAPER_REVISED_3.pdf"
RESULTS = ROOT / "results/v8/gate8/v8-sqrt-error-gate8-20261005-v1"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    checks: list[tuple[str, bool]] = []
    text = MD.read_text(encoding="utf-8")
    lines = text.splitlines()
    new_title = "Cross-Dataset Evaluation and Mechanistic Evidence from Robust Optimization"
    old_title = "Cross-Dataset Confirmation"
    checks.append(("new title present", new_title in text))
    checks.append(("old title absent from latest source", old_title not in text))

    exact = "descriptive 95% paired split-variation interval"
    exact_lines = [line for line in lines if exact in line]
    checks.append(("all exact interval references carry caveat", bool(exact_lines) and all(
        "nominal" in line.lower() and "depend" in line.lower() and "confidence interval" in line.lower()
        for line in exact_lines
    )))
    split_lines = [index for index, line in enumerate(lines) if "95% paired split" in line]
    checks.append(("all 95% paired-split references locally caveated", bool(split_lines) and all(
        "nominal" in " ".join(lines[max(0, index - 3):index + 2]).lower()
        and "confidence" in " ".join(lines[max(0, index - 3):index + 2]).lower()
        for index in split_lines
    )))

    manifest = json.loads((RESULTS / "manifest.json").read_text(encoding="utf-8"))
    config = json.loads((RESULTS / "config.json").read_text(encoding="utf-8"))
    checks.append(("Modal run complete", manifest["status"] == "complete" and manifest["app_id"] == "ap-5Qt6AOVEn68np3HqGt1BFj"))
    checks.append(("40 fits and 40 reused banks", manifest["new_outer_fits"] == 40 and manifest["reused_inner_banks"] == 40))
    checks.append(("all manifest hashes match", all(sha256(RESULTS / name) == expected for name, expected in manifest["sha256"].items())))
    checks.append(("40 immutable checkpoints", len(list((RESULTS / "checkpoints").glob("*.json"))) == 40))
    checks.append(("executed configuration is sqrt/error", config["candidate"]["family"] == "sqrt" and config["candidate"]["selection_metric"] == "error"))
    checks.append(("exact Gate8 registry", config["datasets"] == [
        "blood_transfusion", "mammographicmass_binary", "heart_disease", "parkinson", "wine"
    ] and config["seeds"] == list(range(9000, 9008))))
    checks.append(("frozen reference fingerprint", config["prior_run"]["fingerprint"] == "b6195144643b8d372efa1122e584800a0e1a634582afdc2dbf97823cdeeb99fa"))

    runs = pd.read_csv(RESULTS / "per_run.csv")
    vs_v7 = pd.read_csv(RESULTS / "paired_vs_v7.csv")
    vs_cs = pd.read_csv(RESULTS / "paired_vs_v8_c.csv")
    summary = pd.read_csv(RESULTS / "summary.csv")
    checks.append(("complete paired result tables", len(runs) == 120 and len(vs_v7) == 15 and len(vs_cs) == 15))
    checks.append(("matched model/seed cells", not runs.duplicated(["dataset", "model", "seed"]).any() and set(runs.model) == {"v7", "v8_c", "sqrt_error"}))
    checks.append(("split hashes matched within every cell", all(
        group.train_indices_hash.nunique() == 1 and group.test_indices_hash.nunique() == 1
        for _, group in runs.groupby(["dataset", "seed"])
    )))

    blood = summary[(summary.dataset == "blood_transfusion") & (summary.model == "sqrt_error")].iloc[0]
    checks.append(("Blood Appendix S1 values sourced", all(token in text for token in (
        f"{100 * blood.mean_error:.4f}", f"{100 * blood.balanced_accuracy:.4f}", f"{100 * blood.macro_f1:.4f}"
    ))))
    expected_evidence = {
        ("blood_transfusion", "test_error", "v7"): (-0.008021390374331552, -0.03473893378432989, 0.018696153035666786),
        ("blood_transfusion", "balanced_accuracy", "v7"): (0.06391440132637315, 0.01796931061779166, 0.10985949203495463),
        ("heart_disease", "test_error", "v8_c"): (-0.02333333333333333, -0.05356836842681953, 0.006901701760152864),
        ("heart_disease", "balanced_accuracy", "v8_c"): (-0.010317460317460302, -0.05322279716907126, 0.032587876534150656),
    }
    for (dataset, metric, reference), expected in expected_evidence.items():
        table = vs_v7 if reference == "v7" else vs_cs
        row = table[(table.dataset == dataset) & (table.metric == metric)].iloc[0]
        actual = (row.mean_delta, row.nominal_95_split_interval_low, row.nominal_95_split_interval_high)
        checks.append((f"source values {dataset} {metric} vs {reference}", all(abs(a - b) < 1e-14 for a, b in zip(actual, expected))))

    doc = Document(DOCX)
    checks.append(("DOCX title metadata updated", new_title in doc.core_properties.title and old_title not in doc.core_properties.title))
    checks.append(("DOCX publication objects", len(doc.tables) == 15 and len(doc.inline_shapes) == 5))
    doc_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    checks.append(("DOCX contains targeted Gate8", "Targeted square-root/error-selection Gate8" in doc_text))

    pdf = PdfReader(str(PDF))
    pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    normalized_pdf = re.sub(r"\s+", " ", pdf_text).replace("Cross -Dataset", "Cross-Dataset")
    checks.append(("PDF substantive", len(pdf.pages) >= 18 and len(pdf_text) > 70000))
    checks.append(("PDF title updated", new_title in normalized_pdf and old_title not in normalized_pdf))
    checks.append(("PDF contains targeted Gate8", "Targeted square-root/error-selection Gate8" in normalized_pdf))
    checks.append(("previous version preserved", all((HERE / name).is_file() for name in (
        "SVM_FINAL_PAPER_REVISED_2.md", "SVM_FINAL_PAPER_REVISED_2.docx", "SVM_FINAL_PAPER_REVISED_2.pdf"
    ))))

    failures = [name for name, passed in checks if not passed]
    for name, passed in checks:
        print(f"{'PASS' if passed else 'FAIL'}  {name}")
    if failures:
        raise SystemExit("QA failed: " + "; ".join(failures))
    print(f"ALL {len(checks)} REVISED-3 QA CHECKS PASSED")


if __name__ == "__main__":
    main()
