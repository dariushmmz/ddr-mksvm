"""Build a styled DOCX and PDF from SVM_FINAL_PAPER_REVISED_3.md.

The Markdown manuscript is the canonical content source.  Pandoc performs the
semantic conversion (including equations), python-docx applies journal-style
formatting, and Microsoft Word exports the final PDF without altering content.
"""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "SVM_FINAL_PAPER_REVISED_3.md"
REFERENCE = HERE / ".journal_reference_revised_3.docx"
INTERMEDIATE = HERE / ".SVM_FINAL_PAPER_REVISED_3.pandoc.docx"
DOCX = HERE / "SVM_FINAL_PAPER_REVISED_3.docx"
PDF = HERE / "SVM_FINAL_PAPER_REVISED_3.pdf"

TITLE = (
    "Audited Nonlinear Support Vector Machines for Imbalanced Classification: "
    "Cross-Dataset Evaluation and Mechanistic Evidence from Robust Optimization"
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def set_repeat_xml(paragraph, keep_next: bool = False, keep_lines: bool = False) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if keep_next:
        ppr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        ppr.append(OxmlElement("w:keepLines"))


def _tab_run() -> OxmlElement:
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:tab"))
    return run


def number_display_equations(doc: Document, expected: int = 14) -> None:
    """Expose Markdown equation tags as Word-native right-aligned numbers.

    Pandoc preserves display mathematics as ``m:oMathPara`` but drops LaTeX
    ``\\tag`` labels in DOCX.  Word's standard center/right tab-stop layout
    keeps the equation centered while placing its sequential number at the
    right margin, without rasterizing or changing the mathematics.
    """

    display = [p for p in doc.paragraphs if p._p.xpath("./m:oMathPara")]
    if len(display) != expected:
        raise RuntimeError(
            f"Expected {expected} display equations, found {len(display)}; "
            "refusing to assign potentially incorrect numbers."
        )

    for number, paragraph in enumerate(display, start=1):
        math_para = paragraph._p.xpath("./m:oMathPara")[0]
        math_nodes = math_para.findall(qn("m:oMath"))
        if len(math_nodes) != 1:
            raise RuntimeError(
                f"Equation {number} contains {len(math_nodes)} math nodes; expected one."
            )
        math = math_nodes[0]
        math_para.remove(math)
        location = paragraph._p.index(math_para)
        paragraph._p.remove(math_para)
        paragraph._p.insert(location, _tab_run())
        paragraph._p.insert(location + 1, math)
        paragraph._p.insert(location + 2, _tab_run())
        label = paragraph.add_run(f"({number})")
        label.font.name = "Times New Roman"
        label.font.size = Pt(10)
        label._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

        tabs = paragraph.paragraph_format.tab_stops
        tabs.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)


def create_reference() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(25, 55, 85)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)

    for name, size, before, after in [
        ("Heading 1", 14, 14, 5),
        ("Heading 2", 12, 10, 4),
        ("Heading 3", 11, 8, 3),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(25, 55, 85)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Caption" not in styles:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(45, 45, 45)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True

    if "Compact" not in styles:
        styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
    compact = styles["Compact"]
    compact.font.name = "Times New Roman"
    compact.font.size = Pt(9.5)
    compact.paragraph_format.space_after = Pt(1)
    compact.paragraph_format.line_spacing = 1.05

    doc.add_paragraph("Reference style document")
    doc.save(REFERENCE)


def pandoc_convert() -> None:
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise RuntimeError("Pandoc is required but was not found on PATH")
    cmd = [
        pandoc,
        str(SOURCE),
        "--from=markdown+tex_math_dollars+pipe_tables+link_attributes",
        "--to=docx",
        f"--reference-doc={REFERENCE}",
        f"--resource-path={HERE}",
        "--standalone",
        f"--output={INTERMEDIATE}",
    ]
    subprocess.run(cmd, check=True, cwd=HERE)


def style_docx() -> None:
    doc = Document(INTERMEDIATE)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Nonlinear SVM architecture, class imbalance, and robust optimization"
    doc.core_properties.keywords = (
        "support vector machine; class imbalance; robust optimization; reproducibility; numerical conditioning"
    )
    doc.core_properties.comments = "Generated from SVM_FINAL_PAPER_REVISED_3.md; formatting pass only."

    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2.1)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

        header = section.header
        hp = header.paragraphs[0]
        hp.text = "Audited nonlinear SVMs for imbalanced classification"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in hp.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(8)
            run.font.italic = True
            run.font.color.rgb = RGBColor(90, 90, 90)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prefix = fp.add_run("Page ")
        prefix.font.name = "Times New Roman"
        prefix.font.size = Pt(8)
        add_page_field(fp)

    # Remove conversion-only empty paragraphs while preserving equations,
    # drawings, explicit breaks, and bookmarks. This prevents large vertical
    # gaps (notably the historical Section 5.5 defect) without changing content.
    for p in list(doc.paragraphs):
        if p.text.strip():
            continue
        if p._p.xpath(".//w:drawing | .//m:oMath | .//m:oMathPara | .//w:br | .//w:bookmarkStart"):
            continue
        parent = p._element.getparent()
        if parent is not None:
            parent.remove(p._element)

    # First nonempty paragraph is the manuscript title.
    first_nonempty = next(p for p in doc.paragraphs if p.text.strip())
    first_nonempty.style = doc.styles["Title"]
    set_repeat_xml(first_nonempty, keep_next=True, keep_lines=True)

    in_references = False
    for p in doc.paragraphs:
        text = p.text.strip()
        style_name = p.style.name if p.style else ""
        if text == "References":
            in_references = True
        elif text == "Data and Code Availability":
            in_references = False
        if p is first_nonempty:
            continue
        if text.startswith("Anonymous manuscript draft") or text.startswith("Author names,"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.italic = True
                run.font.color.rgb = RGBColor(70, 70, 70)
        elif style_name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            if text in {"References", "Data and Code Availability", "Declarations"}:
                p.paragraph_format.page_break_before = text == "References"
        elif style_name == "Caption" or text.startswith("Figure "):
            p.style = doc.styles["Caption"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
        elif text.startswith("Table "):
            p.style = doc.styles["Caption"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
        # Pandoc encodes both inline and display mathematics with ``m:oMath``.
        # Only display equations are wrapped in ``m:oMathPara``; centering every
        # paragraph containing inline notation would incorrectly center much of
        # the scientific prose.
        elif p._p.xpath(".//m:oMathPara"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
        elif text:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.widow_control = True
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.space_after = Pt(4)

        if in_references and text != "References":
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.keep_together = True

        # Standardize fonts without overriding bold/italic/math objects.
        for run in p.runs:
            if run._r.xpath(".//m:oMath") or run._r.xpath(".//m:oMathPara"):
                continue
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            if in_references and text != "References":
                run.font.size = Pt(9.2)

    number_display_equations(doc)

    for table in doc.tables:
        table.autofit = True
        table.alignment = 1
        table.style = "Table Grid"
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for ridx, row in enumerate(table.rows):
            prevent_row_split(row)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ridx == 0:
                    set_cell_shading(cell, "D9E7F5")
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.keep_together = True
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                        run.font.size = Pt(7.4 if len(table.columns) >= 7 else 8.0)
                        if ridx == 0:
                            run.font.bold = True

    # Center paragraphs containing only an image and keep them with captions.
    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)

    doc.save(DOCX)


def export_pdf() -> None:
    import win32com.client  # type: ignore

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    document = None
    try:
        document = word.Documents.Open(str(DOCX.resolve()), ReadOnly=True)
        try:
            document.Fields.Update()
            document.Repaginate()
        except Exception:
            pass
        # 17 = wdExportFormatPDF; 0 = all document; 0 = document content.
        document.ExportAsFixedFormat(
            OutputFileName=str(PDF.resolve()),
            ExportFormat=17,
            OpenAfterExport=False,
            OptimizeFor=0,
            Range=0,
            Item=0,
            IncludeDocProps=True,
            KeepIRM=True,
            CreateBookmarks=1,
            DocStructureTags=True,
            BitmapMissingFonts=True,
            UseISO19005_1=False,
        )
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()


def main() -> None:
    create_reference()
    pandoc_convert()
    style_docx()
    export_pdf()
    print(f"DOCX: {DOCX} ({DOCX.stat().st_size} bytes)")
    print(f"PDF:  {PDF} ({PDF.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
